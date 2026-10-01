from __future__ import annotations

import base64
import io
import re
from datetime import datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from fastapi import Cookie, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image, LongTable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .cache_reads import CacheReadError, cache_get
from .config import get_settings


settings = get_settings()
_TZ = ZoneInfo("America/Recife")
ROOT = Path(__file__).resolve().parents[1]
INDUSTRIES_FILE = ROOT / "frontend" / "industries.html"
_INSTALLED = False

_TAB_LABELS = {
    "overview": "Visão geral",
    "critical": "Críticos",
    "rupture": "Risco de ruptura",
    "low": "Venda abaixo da média",
    "high": "Venda acima do normal",
    "dde": "DDE",
    "noTurn": "Sem giro",
    "suppliers": "Fornecedores",
}


def _logo_bytes() -> bytes | None:
    try:
        text = INDUSTRIES_FILE.read_text(encoding="utf-8")
    except Exception:
        return None
    match = re.search(r'<div class="brand"><img src="data:image/png;base64,([^"]+)"', text)
    if not match:
        return None
    try:
        return base64.b64decode(match.group(1))
    except Exception:
        return None


def _status(item: dict[str, Any]) -> str:
    if item.get("produtoNovo"):
        return "PRODUTO NOVO"
    if item.get("ruptura"):
        return "RUPTURA"
    if item.get("riscoRupturaAlto"):
        return "RISCO ALTO DE RUPTURA"
    if item.get("riscoRuptura"):
        return "RISCO DE RUPTURA"
    if item.get("semGiro"):
        return "SEM GIRO"
    if item.get("critico"):
        return "CRÍTICO"
    if item.get("alta"):
        return "EM ALTA"
    if item.get("baixo"):
        return "EM QUEDA"
    return "NORMAL"


def _tab_match(item: dict[str, Any], tab: str) -> bool:
    if tab == "critical":
        return bool(item.get("critico"))
    if tab == "rupture":
        return bool(item.get("riscoRuptura"))
    if tab == "low":
        return bool(item.get("baixo"))
    if tab == "high":
        return bool(item.get("alta"))
    if tab == "dde":
        return float(item.get("estoque") or 0) > 0
    if tab == "noTurn":
        return bool(item.get("semGiro"))
    return True


def _sort_products(rows: list[dict[str, Any]], tab: str, sort_key: str, direction: int) -> list[dict[str, Any]]:
    allowed = {
        "produto",
        "fornecedor",
        "estoque",
        "mediaUnidades",
        "dde",
        "ultimaEntrada",
        "variacaoPct",
        "acao",
        "codigo",
    }
    key = sort_key if sort_key in allowed else "produto"
    direction = -1 if direction < 0 else 1
    if tab == "dde" and sort_key == "produto":
        key = "dde"
        direction = -1

    def sort_value(item: dict[str, Any]):
        value = item.get(key)
        if isinstance(value, (int, float)):
            return float(value)
        return str(value or "").casefold()

    return sorted(rows, key=sort_value, reverse=direction < 0)


def _filter_products(
    result: dict[str, Any],
    *,
    q: str,
    suppliers: list[str],
    curve: str,
    tab: str,
    sort_key: str,
    direction: int,
) -> list[dict[str, Any]]:
    needle = q.strip().casefold()
    supplier_set = {str(value).strip() for value in suppliers if str(value).strip()}
    curve = curve.strip()
    rows: list[dict[str, Any]] = []
    for item in result.get("produtos") or []:
        if not isinstance(item, dict):
            continue
        if supplier_set and str(item.get("fornecedor") or "") not in supplier_set:
            continue
        if curve and str(item.get("curva") or "") != curve:
            continue
        if needle:
            haystack = " ".join(
                str(item.get(key) or "")
                for key in ("produto", "codigo", "ean", "fornecedor")
            ).casefold()
            if needle not in haystack:
                continue
        if not _tab_match(item, tab):
            continue
        rows.append(item)
    return _sort_products(rows, tab, sort_key, direction)


def _filter_suppliers(result: dict[str, Any], *, q: str, suppliers: list[str]) -> list[dict[str, Any]]:
    needle = q.strip().casefold()
    supplier_set = {str(value).strip() for value in suppliers if str(value).strip()}
    rows: list[dict[str, Any]] = []
    for item in result.get("fornecedores") or []:
        if not isinstance(item, dict):
            continue
        name = str(item.get("fornecedor") or "")
        if supplier_set and name not in supplier_set:
            continue
        if needle and needle not in name.casefold():
            continue
        rows.append(item)
    return rows


def _filter_description(q: str, suppliers: list[str], curve: str, tab: str) -> str:
    parts = [f"Visão: {_TAB_LABELS.get(tab, 'Visão geral')}"]
    selected = [str(value).strip() for value in suppliers if str(value).strip()]
    if selected:
        parts.append("Fornecedores: " + ", ".join(selected))
    else:
        parts.append("Fornecedores: todos")
    if curve.strip():
        parts.append("Curva: " + curve.strip())
    if q.strip():
        parts.append("Busca: " + q.strip())
    return " | ".join(parts)


async def _analysis(session: str | None) -> dict[str, Any]:
    from . import commercial_intelligence as ci

    ci._profile(session)
    try:
        payload, meta = await cache_get(modulo="MAPA_ESTOQUE", settings=settings)
    except CacheReadError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    return ci._analyze(payload, meta)


def _excel_bytes(
    result: dict[str, Any],
    rows: list[dict[str, Any]],
    *,
    tab: str,
    description: str,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Inteligência Comercial"
    green = "075B49"
    pale_green = "E8F2EF"

    if tab == "suppliers":
        headers = ["Fornecedor", "Produtos", "Críticos", "Em queda", "Em alta", "Sem giro", "Valor estoque"]
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        ws["A1"] = "DISMEPE ONE — Inteligência Comercial"
        ws["A2"] = description
        ws["A3"] = "Gerado em " + datetime.now(_TZ).strftime("%d/%m/%Y %H:%M")
        header_row = 5
        for col, header in enumerate(headers, 1):
            cell = ws.cell(header_row, col, header)
            cell.fill = PatternFill("solid", fgColor=green)
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
        for r, item in enumerate(rows, header_row + 1):
            values = [
                item.get("fornecedor"), item.get("produtos"), item.get("criticos"),
                item.get("queda"), item.get("alta"), item.get("semGiro"), item.get("valorEstoque"),
            ]
            for c, value in enumerate(values, 1):
                ws.cell(r, c, value)
        widths = [42, 12, 12, 12, 12, 12, 18]
    else:
        months = list(result.get("meses") or [])[:4]
        headers = [
            "Código", "Produto", "Fornecedor", "Curva", "Estoque", "Média", "DDE",
            "Última entrada", "Variação %", "Status", "Ação", *months,
        ]
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        ws["A1"] = "DISMEPE ONE — Inteligência Comercial"
        ws["A2"] = description
        ws["A3"] = "Gerado em " + datetime.now(_TZ).strftime("%d/%m/%Y %H:%M")
        header_row = 5
        for col, header in enumerate(headers, 1):
            cell = ws.cell(header_row, col, header)
            cell.fill = PatternFill("solid", fgColor=green)
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")

        status_fills = {
            "PRODUTO NOVO": PatternFill("solid", fgColor="DBEAFE"),
            "RUPTURA": PatternFill("solid", fgColor="FECACA"),
            "RISCO ALTO DE RUPTURA": PatternFill("solid", fgColor="FECACA"),
            "RISCO DE RUPTURA": PatternFill("solid", fgColor="FEE2E2"),
            "SEM GIRO": PatternFill("solid", fgColor="FEF3C7"),
        }
        for r, item in enumerate(rows, header_row + 1):
            history = list(item.get("historico") or [])[:4]
            status = _status(item)
            values = [
                item.get("codigo"), item.get("produto"), item.get("fornecedor"), item.get("curva"),
                float(item.get("estoque") or 0), float(item.get("mediaUnidades") or 0),
                int(round(float(item.get("dde") or 0))), item.get("ultimaEntrada") or "",
                float(item.get("variacaoPct") or 0), status, item.get("acao") or "", *history,
            ]
            for c, value in enumerate(values, 1):
                ws.cell(r, c, value)
            fill = status_fills.get(status)
            if fill:
                ws.cell(r, 10).fill = fill
        widths = [12, 48, 32, 10, 12, 12, 10, 16, 13, 24, 28] + [12] * len(months)

    ws["A1"].font = Font(size=16, bold=True, color=green)
    ws["A2"].font = Font(size=10, color="4B635D")
    ws["A3"].font = Font(size=9, color="6B7D78")
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{max(header_row, ws.max_row)}"
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row in range(header_row + 1, ws.max_row + 1):
        if row % 2 == 0:
            for cell in ws[row]:
                if cell.fill.fill_type is None:
                    cell.fill = PatternFill("solid", fgColor="F7FAF9")

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _pdf_bytes(
    result: dict[str, Any],
    rows: list[dict[str, Any]],
    *,
    tab: str,
    description: str,
) -> bytes:
    output = io.BytesIO()
    doc = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        leftMargin=9 * mm,
        rightMargin=9 * mm,
        topMargin=9 * mm,
        bottomMargin=10 * mm,
        title="DISMEPE ONE - Inteligência Comercial",
    )
    styles = getSampleStyleSheet()
    normal = ParagraphStyle(
        "ci-normal",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=5.8,
        leading=7,
        textColor=colors.HexColor("#17332C"),
    )
    small = ParagraphStyle(
        "ci-small",
        parent=normal,
        fontSize=5.2,
        leading=6.3,
        textColor=colors.HexColor("#52645F"),
    )
    title_style = ParagraphStyle(
        "ci-title",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=14,
        leading=16,
        textColor=colors.HexColor("#075B49"),
        spaceAfter=2 * mm,
    )

    story: list[Any] = []
    logo = _logo_bytes()
    if logo:
        image = Image(io.BytesIO(logo), width=46 * mm, height=16.4 * mm)
        story.append(image)
    story.append(Paragraph("Inteligência Comercial", title_style))
    story.append(Paragraph(description, small))
    story.append(Paragraph("Gerado em " + datetime.now(_TZ).strftime("%d/%m/%Y %H:%M"), small))
    story.append(Spacer(1, 2.5 * mm))

    if tab == "suppliers":
        data: list[list[Any]] = [[
            "Fornecedor", "Produtos", "Críticos", "Em queda", "Em alta", "Sem giro", "Valor estoque"
        ]]
        for item in rows:
            data.append([
                Paragraph(str(item.get("fornecedor") or ""), normal),
                item.get("produtos") or 0,
                item.get("criticos") or 0,
                item.get("queda") or 0,
                item.get("alta") or 0,
                item.get("semGiro") or 0,
                f"R$ {float(item.get('valorEstoque') or 0):,.2f}".replace(",", "X").replace(".", ",").replace("X", "."),
            ])
        table = LongTable(data, repeatRows=1, colWidths=[72 * mm, 22 * mm, 22 * mm, 22 * mm, 22 * mm, 22 * mm, 30 * mm])
        style = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#075B49")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 6),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D8E5E1")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7FAF9")]),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ]
        table.setStyle(TableStyle(style))
        story.append(table)
    else:
        legend = Table(
            [["Legenda", "Produto novo", "Risco de ruptura", "Sem giro"]],
            colWidths=[22 * mm, 31 * mm, 34 * mm, 25 * mm],
        )
        legend.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 6),
            ("BACKGROUND", (1, 0), (1, 0), colors.HexColor("#DBEAFE")),
            ("BACKGROUND", (2, 0), (2, 0), colors.HexColor("#FEE2E2")),
            ("BACKGROUND", (3, 0), (3, 0), colors.HexColor("#FEF3C7")),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CBD7D3")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        story.append(legend)
        story.append(Spacer(1, 2.5 * mm))

        months = list(result.get("meses") or [])[:4]
        headers = ["Cód.", "Produto", "Fornecedor", "Est.", "Média", "DDE", "Últ. entrada", "Var.%", "Status", *months]
        data: list[list[Any]] = [headers]
        row_backgrounds: list[tuple[int, colors.Color]] = []
        for index, item in enumerate(rows, 1):
            history = list(item.get("historico") or [])[:4]
            status = _status(item)
            data.append([
                str(item.get("codigo") or ""),
                Paragraph(str(item.get("produto") or ""), normal),
                Paragraph(str(item.get("fornecedor") or ""), small),
                f"{float(item.get('estoque') or 0):g}",
                f"{float(item.get('mediaUnidades') or 0):g}",
                str(int(round(float(item.get("dde") or 0)))),
                str(item.get("ultimaEntrada") or ""),
                f"{float(item.get('variacaoPct') or 0):.0f}%",
                Paragraph(status, small),
                *[f"{float(value or 0):g}" for value in history],
            ])
            if item.get("produtoNovo"):
                row_backgrounds.append((index, colors.HexColor("#DBEAFE")))
            elif item.get("riscoRuptura"):
                row_backgrounds.append((index, colors.HexColor("#FEE2E2")))
            elif item.get("semGiro"):
                row_backgrounds.append((index, colors.HexColor("#FEF3C7")))

        table = LongTable(
            data,
            repeatRows=1,
            colWidths=[12 * mm, 52 * mm, 34 * mm, 15 * mm, 15 * mm, 12 * mm, 22 * mm, 14 * mm, 27 * mm] + [14 * mm] * len(months),
        )
        style: list[tuple[Any, ...]] = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#075B49")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 5.6),
            ("GRID", (0, 0), (-1, -1), 0.2, colors.HexColor("#D8E5E1")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 2),
            ("RIGHTPADDING", (0, 0), (-1, -1), 2),
            ("TOPPADDING", (0, 0), (-1, -1), 2.5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
        ]
        for row_index, background in row_backgrounds:
            style.append(("BACKGROUND", (0, row_index), (-1, row_index), background))
        table.setStyle(TableStyle(style))
        story.append(table)

    def footer(canvas, _doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 6.5)
        canvas.setFillColor(colors.HexColor("#60746F"))
        canvas.drawRightString(landscape(A4)[0] - 9 * mm, 5 * mm, f"Página {canvas.getPageNumber()}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def install_commercial_intelligence_exports(app: Any) -> None:
    global _INSTALLED
    if _INSTALLED:
        return
    _INSTALLED = True

    async def export_excel(
        session: str | None = Cookie(default=None, alias=settings.cookie_name),
        q: str = "",
        fornecedor: list[str] | None = Query(default=None),
        curva: str = "",
        aba: str = "overview",
        ordenar: str = "produto",
        direcao: int = 1,
    ):
        result = await _analysis(session)
        suppliers = fornecedor or []
        description = _filter_description(q, suppliers, curva, aba)
        if aba == "suppliers":
            rows = _filter_suppliers(result, q=q, suppliers=suppliers)
        else:
            rows = _filter_products(
                result,
                q=q,
                suppliers=suppliers,
                curve=curva,
                tab=aba,
                sort_key=ordenar,
                direction=direcao,
            )
        body = _excel_bytes(result, rows, tab=aba, description=description)
        filename = "inteligencia_comercial_" + datetime.now(_TZ).strftime("%Y%m%d_%H%M") + ".xlsx"
        return StreamingResponse(
            io.BytesIO(body),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}"', "Cache-Control": "no-store"},
        )

    async def export_pdf(
        session: str | None = Cookie(default=None, alias=settings.cookie_name),
        q: str = "",
        fornecedor: list[str] | None = Query(default=None),
        curva: str = "",
        aba: str = "overview",
        ordenar: str = "produto",
        direcao: int = 1,
    ):
        result = await _analysis(session)
        suppliers = fornecedor or []
        description = _filter_description(q, suppliers, curva, aba)
        if aba == "suppliers":
            rows = _filter_suppliers(result, q=q, suppliers=suppliers)
        else:
            rows = _filter_products(
                result,
                q=q,
                suppliers=suppliers,
                curve=curva,
                tab=aba,
                sort_key=ordenar,
                direction=direcao,
            )
        body = _pdf_bytes(result, rows, tab=aba, description=description)
        filename = "inteligencia_comercial_" + datetime.now(_TZ).strftime("%Y%m%d_%H%M") + ".pdf"
        return StreamingResponse(
            io.BytesIO(body),
            media_type="application/pdf",
            headers={"Content-Disposition": f'attachment; filename="{filename}"', "Cache-Control": "no-store"},
        )

    app.add_api_route("/data/inteligencia-comercial/export.xlsx", export_excel, methods=["GET"])
    app.add_api_route("/data/inteligencia-comercial/export.pdf", export_pdf, methods=["GET"])
