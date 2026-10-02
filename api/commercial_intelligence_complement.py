from __future__ import annotations

import csv
import io
import os
import unicodedata
from datetime import date, datetime
from pathlib import Path
from typing import Any
from urllib.parse import unquote

import httpx
from fastapi import Cookie, HTTPException, Request
from openpyxl import load_workbook

from .config import get_settings
from .security import normalizar


settings = get_settings()
_INSTALLED = False
_MAX_UPLOAD_BYTES = 12 * 1024 * 1024

_CODE_KEYS = {
    "codigo", "cod", "cod produto", "codigo produto", "codigo do produto",
    "cod. produto", "cod produto", "codigo mercadoria", "cod mercadoria",
}
_PRICE_KEYS = {
    "preco medio", "preco medio produto", "custo medio", "preco custo",
    "custo", "preco de custo", "custo medio produto",
}
_LOT_KEYS = {"lote", "lote produto", "numero lote", "nr lote", "n lote"}
_EXPIRY_KEYS = {
    "validade", "vencimento", "data validade", "data de validade",
    "data vencimento", "data de vencimento", "dt validade", "dt vencimento",
}
_QTY_KEYS = {
    "quantidade ultima entrada", "qtd ultima entrada", "qtde ultima entrada",
    "quant ultima entrada", "quantidade da ultima entrada", "qtd da ultima entrada",
    "quantidade entrada", "qtd entrada", "qtde entrada",
}


def _norm(value: Any) -> str:
    raw = str(value or "").strip().lower()
    raw = unicodedata.normalize("NFD", raw)
    raw = "".join(ch for ch in raw if unicodedata.category(ch) != "Mn")
    raw = raw.replace("_", " ").replace("-", " ")
    return " ".join(raw.split())


def _code(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    raw = str(value).strip()
    if raw.endswith(".0") and raw[:-2].replace("-", "").isdigit():
        raw = raw[:-2]
    return raw


def _num(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    raw = str(value).strip().replace("R$", "").replace(" ", "")
    if not raw:
        return 0.0
    if "," in raw:
        raw = raw.replace(".", "").replace(",", ".")
    try:
        return float(raw)
    except ValueError:
        return 0.0


def _date_text(value: Any) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    raw = str(value).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d/%m/%y", "%d-%m-%Y"):
        try:
            return datetime.strptime(raw[:10], fmt).strftime("%d/%m/%Y")
        except ValueError:
            continue
    return raw


def _date_sort(value: str) -> tuple[int, str]:
    raw = str(value or "").strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            dt = datetime.strptime(raw[:10], fmt)
            return (0, dt.strftime("%Y%m%d"))
        except ValueError:
            continue
    return (1, raw)


def _field_key(header: Any) -> str | None:
    h = _norm(header)
    if h in _CODE_KEYS or ("codigo" in h and "produto" in h):
        return "codigo"
    if h in _PRICE_KEYS or (("preco" in h or "custo" in h) and "medio" in h):
        return "precoMedio"
    if h in _LOT_KEYS or h.startswith("lote"):
        return "lote"
    if h in _EXPIRY_KEYS or "validade" in h or "vencimento" in h:
        return "vencimento"
    if h in _QTY_KEYS or (("qtd" in h or "qtde" in h or "quant" in h) and "entrada" in h):
        return "quantidadeUltimaEntrada"
    return None


def _detect_header(rows: list[list[Any]]) -> tuple[int, dict[str, int]]:
    best: tuple[int, dict[str, int]] | None = None
    for idx, row in enumerate(rows[:30]):
        mapping: dict[str, int] = {}
        for col, value in enumerate(row):
            key = _field_key(value)
            if key and key not in mapping:
                mapping[key] = col
        if "codigo" in mapping and len(mapping) >= 2:
            if best is None or len(mapping) > len(best[1]):
                best = (idx, mapping)
    if best is None:
        raise HTTPException(
            status_code=400,
            detail=(
                "Não encontrei a coluna de código junto com Preço Médio, Lote, Validade "
                "ou Quantidade da Última Entrada no arquivo."
            ),
        )
    return best


def _aggregate(raw_rows: list[list[Any]], header_row: int, mapping: dict[str, int]) -> list[dict[str, Any]]:
    grouped: dict[str, dict[str, Any]] = {}
    for row in raw_rows[header_row + 1:]:
        def cell(key: str) -> Any:
            pos = mapping.get(key)
            return row[pos] if pos is not None and pos < len(row) else None

        code = _code(cell("codigo"))
        if not code:
            continue
        price = max(0.0, _num(cell("precoMedio")))
        lot = str(cell("lote") or "").strip()
        expiry = _date_text(cell("vencimento"))
        qty = max(0.0, _num(cell("quantidadeUltimaEntrada")))

        target = grouped.setdefault(code, {
            "codigo": code,
            "precoMedio": 0.0,
            "lote": "",
            "vencimento": "",
            "quantidadeUltimaEntrada": 0.0,
            "lotes": [],
        })
        if price > 0:
            target["precoMedio"] = round(price, 4)
        if lot or expiry or qty > 0:
            target["lotes"].append({
                "lote": lot,
                "vencimento": expiry,
                "quantidadeUltimaEntrada": round(qty, 3),
            })

    for target in grouped.values():
        lots = target.get("lotes") or []
        if lots:
            lots.sort(key=lambda item: _date_sort(str(item.get("vencimento") or "")))
            primary = lots[0]
            target["lote"] = str(primary.get("lote") or "")
            target["vencimento"] = str(primary.get("vencimento") or "")
            target["quantidadeUltimaEntrada"] = float(primary.get("quantidadeUltimaEntrada") or 0)
    return list(grouped.values())


def _parse_xlsx(content: bytes) -> list[dict[str, Any]]:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Não foi possível abrir o arquivo Excel.") from exc

    best: list[dict[str, Any]] = []
    for ws in wb.worksheets:
        rows = [list(row) for row in ws.iter_rows(values_only=True)]
        if not rows:
            continue
        try:
            header_row, mapping = _detect_header(rows)
        except HTTPException:
            continue
        parsed = _aggregate(rows, header_row, mapping)
        if len(parsed) > len(best):
            best = parsed
    if not best:
        raise HTTPException(status_code=400, detail="Nenhuma linha válida foi encontrada no Excel.")
    return best


def _parse_csv(content: bytes) -> list[dict[str, Any]]:
    text = None
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise HTTPException(status_code=400, detail="Não foi possível ler o CSV.")

    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        reader = csv.reader(io.StringIO(text), dialect)
    except csv.Error:
        reader = csv.reader(io.StringIO(text), delimiter=";")
    rows = [list(row) for row in reader]
    header_row, mapping = _detect_header(rows)
    parsed = _aggregate(rows, header_row, mapping)
    if not parsed:
        raise HTTPException(status_code=400, detail="Nenhuma linha válida foi encontrada no CSV.")
    return parsed


def _parse_file(file_name: str, content: bytes) -> list[dict[str, Any]]:
    suffix = Path(file_name).suffix.lower()
    if suffix == ".xlsx":
        return _parse_xlsx(content)
    if suffix == ".csv":
        return _parse_csv(content)
    raise HTTPException(
        status_code=400,
        detail="Neste momento o mapa complementar deve ser enviado em Excel (.xlsx) ou CSV (.csv).",
    )


def _secret() -> str:
    value = str(os.getenv("DISMEPE_COMMERCIAL_COMPLEMENT_SECRET") or "").strip()
    if not value:
        raise HTTPException(status_code=503, detail="Armazenamento do mapa complementar indisponível.")
    return value


def _owner(session: str | None) -> str:
    from . import commercial_intelligence as ci

    profile = ci._profile(session)
    owner = normalizar(profile.get("usuario") or profile.get("sub") or profile.get("nome") or "")
    if not owner:
        raise HTTPException(status_code=401, detail="Usuário da sessão não identificado.")
    return owner


async def _rpc(action: str, owner: str, payload: dict[str, Any] | None = None) -> dict[str, Any]:
    endpoint = settings.supabase_url.rstrip("/") + "/rest/v1/rpc/dismepe_commercial_complement_api"
    key = settings.supabase_publishable_key
    headers = {
        "apikey": key,
        "authorization": f"Bearer {key}",
        "content-type": "application/json",
        "accept": "application/json",
    }
    body = {
        "p_secret": _secret(),
        "p_action": action,
        "p_owner": owner,
        "p_payload": payload or {},
    }
    try:
        async with httpx.AsyncClient(timeout=max(12.0, settings.request_timeout_seconds)) as client:
            response = await client.post(endpoint, headers=headers, json=body)
    except (httpx.TimeoutException, httpx.NetworkError) as exc:
        raise HTTPException(status_code=503, detail="O armazenamento do mapa complementar não respondeu.") from exc

    try:
        data = response.json()
    except ValueError as exc:
        raise HTTPException(status_code=503, detail="Resposta inválida do armazenamento complementar.") from exc
    if response.status_code < 200 or response.status_code >= 300:
        message = str(data.get("message") or data.get("details") or "") if isinstance(data, dict) else ""
        raise HTTPException(status_code=503, detail=message or "Falha ao acessar o mapa complementar.")
    if not isinstance(data, dict):
        raise HTTPException(status_code=503, detail="Formato inválido do mapa complementar.")
    return data


async def merge_complement_for_session(payload: dict[str, Any], session: str | None) -> tuple[dict[str, Any], dict[str, Any]]:
    owner = _owner(session)
    data = await _rpc("GET", owner)
    if data.get("encontrado") is not True:
        return payload, {"ativo": False}
    record = data.get("complemento") if isinstance(data.get("complemento"), dict) else {}
    comp_payload = record.get("payload") if isinstance(record.get("payload"), dict) else {}
    lines = comp_payload.get("linhas") if isinstance(comp_payload.get("linhas"), list) else []
    by_code = {
        str(item.get("codigo") or "").strip(): item
        for item in lines
        if isinstance(item, dict) and str(item.get("codigo") or "").strip()
    }

    matched = 0
    main_rows = payload.get("linhas") if isinstance(payload.get("linhas"), list) else []
    for row in main_rows:
        if not isinstance(row, dict):
            continue
        code = str(row.get("codigo") or "").strip()
        extra = by_code.get(code)
        if not extra:
            continue
        matched += 1
        price = float(extra.get("precoMedio") or 0)
        if price > 0:
            row["preco_mapa_original"] = row.get("preco")
            row["preco"] = price
            row["preco_medio_complemento"] = price
        if extra.get("lote"):
            row["lote"] = extra.get("lote")
        if extra.get("vencimento"):
            row["validade"] = extra.get("vencimento")
            row["vencimento"] = extra.get("vencimento")
        row["quantidade_ultima_entrada"] = float(extra.get("quantidadeUltimaEntrada") or 0)
        row["lotes_complemento"] = extra.get("lotes") if isinstance(extra.get("lotes"), list) else []
        row["complemento_temporario"] = True

    meta = {
        "ativo": True,
        "arquivo": str(record.get("file_name") or ""),
        "enviadoEm": str(record.get("uploaded_at") or ""),
        "codigosComplementares": int(record.get("row_count") or len(lines)),
        "produtosCruzados": matched,
    }
    payload["complementoTemporario"] = meta
    return payload, meta


def install_commercial_intelligence_complement(app: Any) -> None:
    global _INSTALLED
    if _INSTALLED:
        return
    _INSTALLED = True

    from . import commercial_intelligence as ci

    original_product = ci._product

    def product_with_complement(row: dict[str, Any], labels: list[str]) -> dict[str, Any]:
        item = original_product(row, labels)
        item["quantidadeUltimaEntrada"] = ci._round(ci._num(row.get("quantidade_ultima_entrada")), 3)
        item["complementoTemporario"] = bool(row.get("complemento_temporario"))
        item["precoMedioComplemento"] = ci._round(ci._num(row.get("preco_medio_complemento")), 4)
        item["precoOriginalMapa"] = ci._round(ci._num(row.get("preco_mapa_original")), 4)
        item["lotesComplemento"] = row.get("lotes_complemento") if isinstance(row.get("lotes_complemento"), list) else []
        return item

    ci._product = product_with_complement

    async def complement_status(session: str | None = Cookie(default=None, alias=settings.cookie_name)):
        owner = _owner(session)
        data = await _rpc("GET", owner)
        if data.get("encontrado") is not True:
            return {"sucesso": True, "ativo": False}
        record = data.get("complemento") if isinstance(data.get("complemento"), dict) else {}
        return {
            "sucesso": True,
            "ativo": True,
            "arquivo": str(record.get("file_name") or ""),
            "enviadoEm": str(record.get("uploaded_at") or ""),
            "codigos": int(record.get("row_count") or 0),
        }

    async def complement_upload(request: Request, session: str | None = Cookie(default=None, alias=settings.cookie_name)):
        owner = _owner(session)
        file_name = unquote(str(request.headers.get("x-file-name") or "").strip())
        if not file_name:
            raise HTTPException(status_code=400, detail="Nome do arquivo não informado.")
        content = await request.body()
        if not content:
            raise HTTPException(status_code=400, detail="Arquivo vazio.")
        if len(content) > _MAX_UPLOAD_BYTES:
            raise HTTPException(status_code=413, detail="O arquivo complementar deve ter no máximo 12 MB.")
        lines = _parse_file(file_name, content)
        stored = await _rpc("UPSERT", owner, {
            "fileName": file_name[:240],
            "rowCount": len(lines),
            "payload": {"linhas": lines},
        })
        record = stored.get("complemento") if isinstance(stored.get("complemento"), dict) else {}
        return {
            "sucesso": True,
            "arquivo": str(record.get("file_name") or file_name),
            "enviadoEm": str(record.get("uploaded_at") or ""),
            "codigos": len(lines),
            "campos": ["Preço Médio", "Lote", "Validade", "Quantidade da Última Entrada"],
        }

    async def complement_delete(session: str | None = Cookie(default=None, alias=settings.cookie_name)):
        owner = _owner(session)
        await _rpc("DELETE", owner)
        return {"sucesso": True}

    app.add_api_route("/data/inteligencia-comercial/complemento", complement_status, methods=["GET"])
    app.add_api_route("/data/inteligencia-comercial/complemento/upload", complement_upload, methods=["POST"])
    app.add_api_route("/data/inteligencia-comercial/complemento", complement_delete, methods=["DELETE"])
