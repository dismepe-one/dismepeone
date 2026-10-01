from __future__ import annotations

import io
import json
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path
from typing import Any

import jwt
from fastapi import APIRouter, BackgroundTasks, Cookie, HTTPException, Request, Response
from fastapi.responses import FileResponse

from .cache_reads import cache_get
from .config import get_settings
from .security import decode_session_token

router = APIRouter()
settings = get_settings()
ROOT = Path(__file__).resolve().parents[1]
PAGE = ROOT / "frontend" / "crm.html"
CRM_BUILD = "CRM-DEV1-20261001"
CRM_MODULE = "CRM_VENDAS_V1"
STOCK_JSON = ROOT / "data" / "industries" / "mapa_estoque_atual.json"
STOCK_FALLBACK = ROOT / "data" / "industries" / "mapa_estoque_2026-09-16.json"


def _norm(v: Any) -> str:
    return re.sub(r"\s+", " ", unicodedata.normalize("NFD", str(v or ""))
                  .encode("ascii", "ignore").decode("ascii").upper()).strip()


def _admin_danton(session: str | None) -> dict[str, Any]:
    if not session:
        raise HTTPException(401, "Sessão ausente.")
    try:
        profile = decode_session_token(session, secret=settings.jwt_secret, issuer=settings.jwt_issuer)
    except jwt.ExpiredSignatureError as exc:
        raise HTTPException(401, "Sessão expirada.") from exc
    except jwt.PyJWTError as exc:
        raise HTTPException(401, "Sessão inválida.") from exc
    tipo = _norm(profile.get("tipo"))
    usuario = _norm(profile.get("usuario") or profile.get("sub"))
    if tipo not in {"ADMIN", "ADMINISTRADOR"} or usuario != "DANTON":
        raise HTTPException(403, "CRM em desenvolvimento: acesso exclusivo do administrador Danton.")
    return profile


def _json(data: dict[str, Any], status: int = 200) -> Response:
    return Response(
        content=json.dumps(data, ensure_ascii=False, separators=(",", ":")),
        status_code=status,
        media_type="application/json",
        headers={"Cache-Control": "no-store, private", "X-DISMEPE-CRM": CRM_BUILD},
    )


async def _edge(acao: str, payload: dict[str, Any]) -> dict[str, Any]:
    import httpx
    url = settings.supabase_url.rstrip("/") + "/functions/v1/dismepe-admin"
    async with httpx.AsyncClient(timeout=httpx.Timeout(60.0)) as client:
        response = await client.post(
            url,
            json={"acao": acao, **payload},
            headers={
                "apikey": settings.supabase_publishable_key,
                "x-dismepe-token": settings.edge_token,
                "Content-Type": "application/json",
                "Accept": "application/json",
            },
        )
    try:
        data = response.json()
    except ValueError as exc:
        raise RuntimeError("Serviço de dados não retornou JSON válido.") from exc
    if not response.is_success or not isinstance(data, dict) or data.get("sucesso") is not True:
        raise RuntimeError(str(data.get("erro") if isinstance(data, dict) else "Falha no serviço de dados."))
    return data


def _stock_map() -> dict[str, dict[str, Any]]:
    for path in (STOCK_JSON, STOCK_FALLBACK):
        try:
            obj = json.loads(path.read_text(encoding="utf-8"))
            rows = obj.get("linhas") if isinstance(obj, dict) else []
            if isinstance(rows, list):
                return {re.sub(r"\D", "", str(r.get("codigo") or "")): r for r in rows if r.get("codigo")}
        except Exception:
            continue
    return {}


async def _positivacao_owner_map() -> dict[str, dict[str, Any]]:
    try:
        payload, _ = await cache_get(modulo="POSITIVACAO_GERAL_V1", settings=settings)
        if payload.get("schema") != "POSITIVACAO_GERAL_V1" or not payload.get("zip"):
            return {}
        import base64, zlib
        raw = zlib.decompress(base64.b64decode(payload["zip"], validate=True))
        data = json.loads(raw)
        rows = data.get("clientes") if isinstance(data, dict) else []
        return {str(r.get("codigo")): r for r in rows if isinstance(r, dict) and r.get("codigo")}
    except Exception:
        return {}


def _row_to_db(row: dict[str, Any], source: str) -> dict[str, Any]:
    def num(v, default=0):
        try:
            return float(v)
        except Exception:
            return default
    data = row.get("Data")
    if isinstance(data, datetime):
        d = data.date().isoformat()
    elif isinstance(data, date):
        d = data.isoformat()
    else:
        d = str(data or "").strip()[:10]
        if "/" in d:
            try:
                d = datetime.strptime(d, "%d/%m/%Y").date().isoformat()
            except ValueError:
                pass
    nf = int(float(row.get("Número NF") or 0))
    cliente = int(float(row.get("Cód. Cliente") or 0))
    produto = int(float(row.get("Cód. Produto") or 0))
    return {
        "data_venda": d,
        "numero_nf": nf,
        "cod_cliente": cliente,
        "cod_produto": produto,
        "total_unidade": num(row.get("Total Unidade")),
        "venda_liquida": num(row.get("Venda Líquida (R$)")),
        "fornecedor": str(row.get("Fornecedor") or "").strip(),
        "chave_origem": f"{d}|{nf}|{cliente}|{produto}|{num(row.get('Total Unidade'))}|{num(row.get('Venda Líquida (R$)'))}",
        "origem_arquivo": source,
    }


@router.get("/crm", include_in_schema=False)
async def crm_page(session: str | None = Cookie(default=None, alias=settings.cookie_name)):
    _admin_danton(session)
    return FileResponse(PAGE, media_type="text/html", headers={"Cache-Control": "no-store, private"})


@router.get("/crm/api/acesso")
async def crm_access(session: str | None = Cookie(default=None, alias=settings.cookie_name)):
    _admin_danton(session)
    return _json({"sucesso": True, "usuario": "DANTON", "build": CRM_BUILD})


@router.get("/crm/api/resumo")
async def crm_summary(session: str | None = Cookie(default=None, alias=settings.cookie_name)):
    _admin_danton(session)
    result = await _edge("CRM_RESUMO", {})
    result["build"] = CRM_BUILD
    return _json(result)


@router.get("/crm/api/produtos")
async def crm_products(session: str | None = Cookie(default=None, alias=settings.cookie_name)):
    _admin_danton(session)
    result = await _edge("CRM_PRODUTOS", {})
    result["build"] = CRM_BUILD
    return _json(result)


@router.get("/crm/api/clientes")
async def crm_clients(session: str | None = Cookie(default=None, alias=settings.cookie_name)):
    _admin_danton(session)
    result = await _edge("CRM_CLIENTES", {})
    result["build"] = CRM_BUILD
    return _json(result)


async def _crm_send_batch_with_retry(batch: list[dict[str, Any]], importacao_id: str) -> None:
    last: Exception | None = None
    for tentativa in range(3):
        try:
            await _edge("CRM_VENDAS_IMPORTAR", {
                "registros": batch,
                "importacao_id": importacao_id,
            })
            return
        except Exception as exc:
            last = exc
            if tentativa < 2:
                import asyncio
                await asyncio.sleep(2 * (tentativa + 1))
    raise last or RuntimeError("Falha ao importar lote CRM.")


async def _crm_process_import(raw: bytes, filename: str, importacao_id: str, total_linhas: int) -> None:
    try:
        ext = "csv" if filename.lower().endswith(".csv") else "xlsx"
        source = f"upload_{datetime.now().strftime('%Y%m%d_%H%M%S')}_{importacao_id[:8]}.{ext}"
        if ext == "csv":
            import csv
            text_csv = raw.decode("utf-8-sig", errors="replace")
            reader = csv.DictReader(io.StringIO(text_csv))
            headers = [str(x or "").strip() for x in (reader.fieldnames or [])]
            expected = ["Data", "Número NF", "Cód. Cliente", "Total Unidade", "Venda Líquida (R$)", "Fornecedor", "Cód. Produto"]
            if headers[:7] != expected:
                raise ValueError("Cabeçalhos incompatíveis. Esperado: Data, Número NF, Cód. Cliente, Total Unidade, Venda Líquida (R$), Fornecedor, Cód. Produto.")
            rows = reader
        else:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            ws = wb.active
            expected = ["Data", "Número NF", "Cód. Cliente", "Total Unidade", "Venda Líquida (R$)", "Fornecedor", "Cód. Produto"]
            rows = (dict(zip(expected, values)) for values in ws.iter_rows(min_row=2, values_only=True))
        batch: list[dict[str, Any]] = []
        processadas = 0
        for row in rows:
            if not row or all(v is None or str(v).strip() == "" for v in row.values()):
                continue
            dbrow = _row_to_db(row, source)
            if dbrow["cod_cliente"] <= 0 or dbrow["cod_produto"] <= 0 or not dbrow["data_venda"]:
                continue
            batch.append(dbrow)
            if len(batch) >= 1000:
                await _crm_send_batch_with_retry(batch, importacao_id)
                processadas += len(batch)
                await _edge("CRM_IMPORT_ATUALIZAR", {"importacao_id": importacao_id, "processadas": processadas, "total_linhas": total_linhas})
                batch = []
        if batch:
            await _crm_send_batch_with_retry(batch, importacao_id)
            processadas += len(batch)
            await _edge("CRM_IMPORT_ATUALIZAR", {"importacao_id": importacao_id, "processadas": processadas, "total_linhas": total_linhas})
        await _edge("CRM_IMPORT_FINALIZAR", {"importacao_id": importacao_id, "sucesso": True})
    except Exception as exc:
        try:
            await _edge("CRM_IMPORT_FINALIZAR", {"importacao_id": importacao_id, "sucesso": False, "erro": str(exc)[:1000]})
        except Exception:
            pass


@router.post("/crm/api/importar")
async def crm_importar(
    request: Request,
    background_tasks: BackgroundTasks,
    session: str | None = Cookie(default=None, alias=settings.cookie_name),
):
    profile = _admin_danton(session)
    raw = await request.body()
    if len(raw) > 35 * 1024 * 1024:
        raise HTTPException(413, "Arquivo acima do limite de 35 MB.")
    if not raw:
        raise HTTPException(400, "Arquivo vazio.")
    filename = request.headers.get("x-filename") or "base_crm"
    filename = filename.replace("%20", " ").strip()
    if not (filename.lower().endswith(".xlsx") or filename.lower().endswith(".csv")):
        raise HTTPException(400, "Formato não suportado. Envie XLSX ou CSV.")
    if filename.lower().endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            total_linhas = max(0, int(wb.active.max_row or 0) - 1)
            wb.close()
        except Exception as exc:
            raise HTTPException(400, f"Não foi possível ler o XLSX: {str(exc)[:300]}") from exc
    else:
        total_linhas = max(0, raw.count(b"\n") - 1)
    created = await _edge("CRM_IMPORT_CRIAR", {
        "nome_arquivo": filename,
        "total_linhas": total_linhas,
        "usuario": str(profile.get("usuario") or "DANTON"),
    })
    importacao_id = str(created.get("id") or "")
    if not importacao_id:
        raise HTTPException(500, "O servidor não retornou o identificador da importação.")
    background_tasks.add_task(_crm_process_import, raw, filename, importacao_id, total_linhas)
    return _json({
        "sucesso": True,
        "aceito": True,
        "importacaoId": importacao_id,
        "totalLinhas": total_linhas,
        "mensagem": "Importação iniciada. O processamento continuará em segundo plano.",
    }, status=202)


@router.get("/crm/api/import/status/{importacao_id}")
async def crm_import_status(
    importacao_id: str,
    session: str | None = Cookie(default=None, alias=settings.cookie_name),
):
    _admin_danton(session)
    return _json(await _edge("CRM_IMPORT_STATUS", {"importacao_id": importacao_id}))

