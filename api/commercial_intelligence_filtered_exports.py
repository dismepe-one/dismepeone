from __future__ import annotations

import copy
import io
import sys
from datetime import datetime
from typing import Any
from zoneinfo import ZoneInfo

from fastapi import Cookie, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image, LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

from .cache_reads import cache_get
from .config import get_settings


settings = get_settings()
_TZ = ZoneInfo("America/Recife")
_INSTALLED = False


def _status(item: dict[str, Any]) -> str:
    if item.get("produtoNovo"):
        return "PRODUTO NOVO"
    if item.get("ruptura"):
        return "RUPTURA"
    if item.get("riscoRupturaAlto"):
        return "RISCO ALTO DE RUPTURA"
    if item.get("riscoRuptura"):
        return "RISCO DE RUPTURA"
    if item.get("semGiro"):
        return "SEM GIRO"
    if item.get("critico"):
        return "CRÍTICO"
    if item.get("alta"):
        return "EM ALTA"
    if item.get("baixo"):
        return "EM QUEDA"
    return "NORMAL"


async def _analysis(session: str | None) -> dict[str, Any]:
    from . import commercial_intelligence as ci
    from .commercial_intelligence_complement import merge_complement_for_session

    ci._profile(session)
    payload, meta = await cache_get(modulo="MAPA_ESTOQUE", settings=settings)
    # O complemento é mesclado em uma cópia para não alterar o snapshot em cache.
    merged = copy.deepcopy(payload)
    merged, _ = await merge_complement_for_session(merged, session)
    return ci._analyze(merged, meta)


def _filter_rows(
    result: dict[str, Any],
    *,
    q: str,
    suppliers: list[str],
    curve: str,
    tab: str,
    sort_key: str,
    direction: int,
    blocked: bool,
    expiring: bool,
) -> list[dict[str, Any]]:
    from . import commercial_intelligence_exports as exports

    rows = exports._filter_products(
        result,
        q=q,
        suppliers=suppliers,
        curve=curve,
        tab=tab,
        sort_key=sort_key,
        direction=direction,
    )
    if blocked:
        rows = [item for item in rows if str(item.get("bloqCompra") or "").strip()]
    if expiring:
        rows = [item for item in rows if bool(item.get("vencimentoProximo"))]
    return rows


def _description(
    *,
    q: str,
    suppliers: list[str],
    curve: str,
    tab: str,
    blocked: bool,
    expiring: bool,
) -> str:
    from . import commercial_intelligence_exports as exports

    base = exports._filter_description(q, suppliers, curve, tab)
    extras: list[str] = []
    if blocked:
        extras.append("Bloq. compra")
    if expiring:
        extras.append("Venc. ≤ 12 meses")
    return base + (" | Filtros: " + " + ".join(extras) if extras else "")


def _excel_bytes(rows: list[dict[str, Any]], description: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Inteligência Comercial"
    headers = [
        "Código", "Produto", "Fornecedor", "Curva", "Bloq. compra",
        "Lote", "Validade", "Qtd. últ. entrada", "Preço médio",
        "Estoque", "Média", "DDE", "Últ. entrada", "Variação %",
        "Status", "Ação",
    ]
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws["A1"] = "DISMEPE ONE — Inteligência Comercial"
    ws["A2"] = description
    ws["A3"] = "Gerado em " + datetime.now(_TZ).strftime("%d/%m/%Y %H:%M")
    ws["A1"].font = Font(size=16, bold=True, color="075B49")
    ws["A2"].font = Font(size=10, color="4B635D")
    ws["A3"].font = Font(size=9, color="6B7D78")

    header_row = 5
    for col, header in enumerate(headers, 1):
        cell = ws.cell(header_row, col, header)
        cell.fill = PatternFill("solid", fgColor="075B49")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_index, item in enumerate(rows, header_row + 1):
        values = [
            str(item.get("codigo") or ""),
            str(item.get("produto") or ""),
            str(item.get("fornecedor") or ""),
            str(item.get("curva") or ""),
            str(item.get("bloqCompra") or ""),
            str(item.get("lote") or ""),
            str(item.get("vencimento") or ""),
            int(round(float(item.get("quantidadeUltimaEntrada") or 0))),
            float(item.get("preco") or 0),
            float(item.get("estoque") or 0),
            float(item.get("mediaUnidades") or 0),
            int(round(float(item.get("dde") or 0))),
            str(item.get("ultimaEntrada") or ""),
            float(item.get("variacaoPct") or 0),
            _status(item),
            str(item.get("acao") or ""),
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row_index, col, value)
        ws.cell(row_index, 9).number_format = 'R$ #,##0.00'
        ws.cell(row_index, 14).number_format = '0.0"%"'
        if item.get("vencimentoProximo"):
            ws.cell(row_index, 7).font = Font(color="BE123C", bold=True)

    widths = [12, 48, 32, 10, 14, 18, 14, 18, 16, 12, 12, 10, 15, 13, 24, 30]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A6"
    ws.auto_filter.ref = f"A5:{get_column_letter(len(headers))}{max(5, ws.max_row)}"
    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def _pdf_bytes(rows: list[dict[str, Any]], description: str) -> bytes:
    from . import commercial_intelligence_exports as exports

    output = io.BytesIO()
    doc = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        leftMargin=8 * mm,
        rightMargin=8 * mm,
        topMargin=8 * mm,
        bottomMargin=9 * mm,
        title="DISMEPE ONE - Inteligência Comercial",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "filtered-title",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=14,
        leading=16,
        textColor=colors.HexColor("#075B49"),
        spaceAfter=2 * mm,
    )
    cell = ParagraphStyle(
        "filtered-cell",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=5.4,
        leading=6.4,
        textColor=colors.HexColor("#17332C"),
    )
    small = ParagraphStyle(
        "filtered-small",
        parent=cell,
        fontSize=5.2,
        leading=6.2,
        textColor=colors.HexColor("#60746F"),
    )

    story: list[Any] = []
    logo = exports._logo_bytes()
    if logo:
        story.append(Image(io.BytesIO(logo), width=44 * mm, height=15.7 * mm))
    story.append(Paragraph("Inteligência Comercial — Recorte filtrado", title_style))
    story.append(Paragraph(description, small))
    story.append(Paragraph("Gerado em " + datetime.now(_TZ).strftime("%d/%m/%Y %H:%M"), small))
    story.append(Spacer(1, 2.5 * mm))

    data: list[list[Any]] = [[
        "Cód.", "Produto", "Fornecedor", "Bloq.", "Lote", "Validade",
        "Qtd. últ.", "Estoque", "DDE", "Status",
    ]]
    for item in rows:
        data.append([
            str(item.get("codigo") or ""),
            Paragraph(str(item.get("produto") or ""), cell),
            Paragraph(str(item.get("fornecedor") or ""), small),
            str(item.get("bloqCompra") or ""),
            str(item.get("lote") or ""),
            str(item.get("vencimento") or ""),
            str(int(round(float(item.get("quantidadeUltimaEntrada") or 0)))),
            f"{float(item.get('estoque') or 0):g}",
            str(int(round(float(item.get("dde") or 0)))),
            Paragraph(_status(item), small),
        ])

    table = LongTable(
        data,
        repeatRows=1,
        colWidths=[12 * mm, 54 * mm, 38 * mm, 18 * mm, 22 * mm, 20 * mm, 20 * mm, 18 * mm, 13 * mm, 29 * mm],
    )
    style: list[tuple[Any, ...]] = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#075B49")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 5.5),
        ("GRID", (0, 0), (-1, -1), 0.2, colors.HexColor("#D8E5E1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 2),
        ("RIGHTPADDING", (0, 0), (-1, -1), 2),
        ("TOPPADDING", (0, 0), (-1, -1), 2.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7FAF9")]),
    ]
    for row_index, item in enumerate(rows, 1):
        if item.get("vencimentoProximo"):
            style.append(("TEXTCOLOR", (5, row_index), (5, row_index), colors.HexColor("#BE123C")))
            style.append(("FONTNAME", (5, row_index), (5, row_index), "Helvetica-Bold"))
    table.setStyle(TableStyle(style))
    story.append(table)
    doc.build(story)
    return output.getvalue()


def install_commercial_intelligence_filtered_exports() -> None:
    global _INSTALLED
    if _INSTALLED:
        return

    main_module = sys.modules.get("api.main")
    app = getattr(main_module, "app", None)
    if app is None:
        return

    async def export_excel(
        session: str | None = Cookie(default=None, alias=settings.cookie_name),
        q: str = "",
        fornecedor: list[str] | None = Query(default=None),
        curva: str = "",
        aba: str = "overview",
        ordenar: str = "produto",
        direcao: int = 1,
        bloq_compra: bool = False,
        vencimento: bool = False,
    ):
        result = await _analysis(session)
        suppliers = fornecedor or []
        rows = _filter_rows(
            result,
            q=q,
            suppliers=suppliers,
            curve=curva,
            tab=aba,
            sort_key=ordenar,
            direction=direcao,
            blocked=bloq_compra,
            expiring=vencimento,
        )
        description = _description(
            q=q,
            suppliers=suppliers,
            curve=curva,
            tab=aba,
            blocked=bloq_compra,
            expiring=vencimento,
        )
        body = _excel_bytes(rows, description)
        filename = "inteligencia_comercial_filtrada_" + datetime.now(_TZ).strftime("%Y%m%d_%H%M") + ".xlsx"
        return StreamingResponse(
            io.BytesIO(body),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}"', "Cache-Control": "no-store"},
        )

    async def export_pdf(
        session: str | None = Cookie(default=None, alias=settings.cookie_name),
        q: str = "",
        fornecedor: list[str] | None = Query(default=None),
        curva: str = "",
        aba: str = "overview",
        ordenar: str = "produto",
        direcao: int = 1,
        bloq_compra: bool = False,
        vencimento: bool = False,
    ):
        result = await _analysis(session)
        suppliers = fornecedor or []
        rows = _filter_rows(
            result,
            q=q,
            suppliers=suppliers,
            curve=curva,
            tab=aba,
            sort_key=ordenar,
            direction=direcao,
            blocked=bloq_compra,
            expiring=vencimento,
        )
        description = _description(
            q=q,
            suppliers=suppliers,
            curve=curva,
            tab=aba,
            blocked=bloq_compra,
            expiring=vencimento,
        )
        body = _pdf_bytes(rows, description)
        filename = "inteligencia_comercial_filtrada_" + datetime.now(_TZ).strftime("%Y%m%d_%H%M") + ".pdf"
        return StreamingResponse(
            io.BytesIO(body),
            media_type="application/pdf",
            headers={"Content-Disposition": f'attachment; filename="{filename}"', "Cache-Control": "no-store"},
        )

    app.add_api_route("/data/inteligencia-comercial/filtered-export.xlsx", export_excel, methods=["GET"])
    app.add_api_route("/data/inteligencia-comercial/filtered-export.pdf", export_pdf, methods=["GET"])
    _INSTALLED = True
