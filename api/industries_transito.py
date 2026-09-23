"""Portal Indústrias — Trânsito. Apenas itens vinculados por CNPJ a laboratório autorizado."""
from __future__ import annotations

import io
import json
import re
import zipfile
from datetime import date, datetime, timedelta
from functools import lru_cache
from zoneinfo import ZoneInfo
from decimal import Decimal, InvalidOperation
from xml.etree import ElementTree as ET

from fastapi import APIRouter, Cookie, HTTPException, Query, Request
from fastapi.responses import Response, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .cache_reads import CacheReadError, cache_get
from .industries import (
    ALL_LABS_VALUE, _buyer_all_labs, _choose_lab, _industry_profile, _is_all_labs_request,
    _portal_lab_key, _strict_admin_profile, _edge_admin_write, settings,
)

router = APIRouter()
MODULE = "INDUSTRIAS_TRANSITO_V1"
SCHEMA = "transito-v1"

# Vinculação EXATA do CNPJ do emitente, conferida no ZIP recebido. Não atribuir
# CNPJ de farmácias/distribuidores a laboratórios por aproximação de nomes.
EMITENTE_LAB = {
    "00677858000195": "ARTE NATIVA",
    "02456955000183": "NATULAB",
    "02456955000507": "NATULAB",
    "02625651000100": "SANFARMA",
    "17115437000173": "GLOBO",
    "03485572000104": "GEOLAB",
    "35356799000138": "LAPON",
    "57235426000303": "BRG SUPLEMENTOS NUTRICIONAIS LTDA",
}

def _tag(element: ET.Element | None, tag: str) -> str:
    if element is None:
        return ""
    child = element.find("{*}" + tag)
    return str(child.text or "").strip() if child is not None else ""

def _xml_root(blob: bytes) -> ET.Element:
    if b"<!DOCTYPE" in blob.upper() or b"<!ENTITY" in blob.upper():
        raise ValueError("DTD e entidades não são permitidas.")
    try:
        return ET.fromstring(blob)
    except ET.ParseError:
        # Alguns arquivos fornecidos são Latin-1 mesmo sem declaração correta.
        # Escapar apenas & isolado, preservando referências XML válidas.
        try:
            content = blob.decode("utf-8")
        except UnicodeDecodeError:
            content = blob.decode("latin-1")
        content = re.sub(r"&(?!#(?:\d+|x[0-9a-fA-F]+);|(?:amp|lt|gt|apos|quot);)", "&amp;", content)
        return ET.fromstring(content)

def _decode_zip(blob: bytes) -> tuple[list[dict], dict]:
    if len(blob) > 10 * 1024 * 1024:
        raise HTTPException(413, "ZIP excede o limite de 10 MB.")
    try:
        archive = zipfile.ZipFile(io.BytesIO(blob))
    except (ValueError, zipfile.BadZipFile) as exc:
        raise HTTPException(400, "Envie um arquivo ZIP válido contendo XMLs de NF-e.") from exc
    names = [entry for entry in archive.infolist() if entry.filename.lower().endswith(".xml") and not entry.is_dir()]
    if not names or len(names) > 2000 or sum(entry.file_size for entry in names) > 40 * 1024 * 1024:
        raise HTTPException(400, "ZIP vazio ou acima do limite seguro de arquivos XML.")
    rows: list[dict] = []
    seen: set[tuple[str, str]] = set()
    errors = 0
    unmapped: set[str] = set()
    ignored = 0
    for entry in names:
        if entry.file_size > 2 * 1024 * 1024:
            errors += 1
            continue
        try:
            root = _xml_root(archive.read(entry))
            nfe = root.find(".//{*}infNFe")
            if nfe is None and root.tag.rsplit("}", 1)[-1] == "infNFe":
                nfe = root
            if nfe is None:
                errors += 1
                continue
            ide = nfe.find("{*}ide")
            emit = nfe.find("{*}emit")
            doc = re.sub(r"\D", "", _tag(emit, "CNPJ"))
            lab = EMITENTE_LAB.get(doc, "")
            if not lab:
                unmapped.add(doc or "SEM CNPJ")
                continue
            # Exclui devoluções, entradas e notas complementares: não são carga
            # de fornecedor em trânsito. Dados antigos não comprovam recebimento.
            if _tag(ide, "tpNF") != "1" or _tag(ide, "finNFe") not in {"", "1"}:
                ignored += 1
                continue
            nat = _tag(ide, "natOp").upper()
            if "DEVOLU" in nat:
                ignored += 1
                continue
            emitted = _tag(ide, "dhEmi") or _tag(ide, "dEmi")
            day = emitted[:10]
            datetime.strptime(day, "%Y-%m-%d")
            key = str(nfe.attrib.get("Id") or "").strip() or entry.filename
            invoice_number = _tag(ide, "nNF")
            if not re.fullmatch(r"\d{1,9}", invoice_number):
                invoice_number = _invoice_number_from_key(key)
            for item in nfe.findall("{*}det"):
                prod = item.find("{*}prod")
                if prod is None:
                    continue
                number = str(item.attrib.get("nItem") or len(rows) + 1)
                unique = (key, number)
                if unique in seen:
                    continue
                seen.add(unique)
                amount = Decimal(_tag(prod, "qCom") or "0")
                if not amount.is_finite() or amount <= 0:
                    continue
                rows.append({
                    "dataEmissao": day,
                    "numeroNFe": invoice_number,
                    "emitente": _tag(emit, "xNome")[:180],
                    "ean": _tag(prod, "cEAN")[:20],
                    "produto": _tag(prod, "xProd")[:260],
                    "quantidade": str(amount),
                    "laboratorio": lab,
                    "chaveItem": key + ":" + number,
                })
        except (ET.ParseError, UnicodeError, ValueError, InvalidOperation, OverflowError):
            errors += 1
    if not rows:
        raise HTTPException(422, "Nenhum item elegível com emitente vinculado a laboratório. A base anterior foi preservada.")
    rows.sort(key=lambda x: (x["dataEmissao"], x["emitente"], x["chaveItem"]), reverse=True)
    return rows, {"arquivos": len(names), "itens": len(rows), "arquivosInvalidos": errors,
                  "naoVinculados": len(unmapped), "operacoesIgnoradas": ignored}


@lru_cache(maxsize=24)
def _national_holidays(year: int) -> frozenset[date]:
    """Feriados nacionais brasileiros e Sexta-feira Santa; não inclui pontos facultativos ou feriados locais."""
    # Computus gregoriano, para posicionar a Sexta-feira Santa a cada ano.
    a, b, c = year % 19, year // 100, year % 100
    d, e = b // 4, b % 4
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i, k = c // 4, c % 4
    l = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l) // 451
    month = (h + l - 7 * m + 114) // 31
    day = (h + l - 7 * m + 114) % 31 + 1
    good_friday = date(year, month, day) - timedelta(days=2)
    return frozenset({
        date(year, 1, 1), date(year, 4, 21), date(year, 5, 1),
        date(year, 9, 7), date(year, 10, 12), date(year, 11, 2),
        date(year, 11, 15), date(year, 11, 20), date(year, 12, 25),
        good_friday,
    })


def _expected_delivery(emission: str) -> date:
    """Quinze dias corridos a partir da emissão, prorrogados até o próximo dia útil."""
    due = date.fromisoformat(emission) + timedelta(days=15)
    while due.weekday() >= 5 or due in _national_holidays(due.year):
        due += timedelta(days=1)
    return due



def _invoice_number_from_key(chave_item: str) -> str:
    """Recupera NF-e de snapshots anteriores somente se a chave possuir 44 dígitos."""
    match = re.fullmatch(r"NFe(\d{44})(?::\d+)?", str(chave_item or "").strip())
    if not match:
        return ""
    return str(int(match.group(1)[25:34]))


def _scope(rows: list[dict], selected: str, *, today: date | None = None) -> list[dict]:
    key = _portal_lab_key(selected)
    filtered = rows if selected == ALL_LABS_VALUE else [
        row for row in rows if _portal_lab_key(row.get("laboratorio")) == key
    ]
    # Data atual de Pernambuco, não o fuso UTC do servidor. O prazo só vence no
    # dia seguinte à previsão; a situação é recalculada a cada consulta.
    reference_day = today if today is not None else datetime.now(ZoneInfo("America/Recife")).date()
    output: list[dict] = []
    for row in filtered:
        due = _expected_delivery(row["dataEmissao"])
        # Nunca devolver identificadores internos ou linhas de outros laboratórios.
        output.append({
            "dataEmissao": row.get("dataEmissao", ""),
            "numeroNFe": str(row.get("numeroNFe") or _invoice_number_from_key(row.get("chaveItem", ""))),
            "previsaoChegada": due.isoformat(),
            "atrasado": reference_day > due,
            "emitente": row.get("emitente", ""),
            "ean": row.get("ean", ""),
            "produto": row.get("produto", ""),
            "quantidade": row.get("quantidade", ""),
        })
    return output

@router.get("/industrias/transito")
async def transit_list(
    laboratorio: str | None = Query(default=None),
    session: str | None = Cookie(default=None, alias=settings.cookie_name),
):
    profile = await _industry_profile(session, require_password_changed=True)
    selected = _choose_lab(profile, laboratorio)
    if _is_all_labs_request(selected) and not _buyer_all_labs(profile):
        raise HTTPException(403, "Visão consolidada não autorizada.")
    try:
        payload, _ = await cache_get(modulo=MODULE, settings=settings)
    except CacheReadError as exc:
        if "ainda não está disponível" in str(exc):
            return Response(content=json.dumps({"sucesso": True, "linhas": [], "importado": False}),
                            media_type="application/json", headers={"Cache-Control": "no-store"})
        raise HTTPException(503, "Não foi possível consultar a base de Trânsito.") from exc
    if payload.get("schema") != SCHEMA or not isinstance(payload.get("linhas"), list):
        raise HTTPException(503, "Fotografia de Trânsito inválida.")
    output = {"sucesso": True, "linhas": _scope(payload["linhas"], selected),
              "importado": True, "atualizadoEm": payload.get("atualizadoEm", "")}
    return Response(content=json.dumps(output, ensure_ascii=False), media_type="application/json",
                    headers={"Cache-Control": "no-store"})


def _excel_text(value: object) -> str:
    """Impedir execução de fórmulas em campos textuais vindos do XML."""
    content = str(value or "").strip()
    return "'" + content if content and content[0] in "=+-@\t\r\n" else content


def _build_transit_excel(rows: list[dict]) -> bytes:
    """Planilha da mesma carteira e das sete colunas mostradas na tela."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Trânsito"
    sheet.append(("DATA DE EMISSÃO", "Nº NF-e", "PREVISÃO DE CHEGADA", "EMITENTE", "EAN", "PRODUTO", "QUANTIDADE"))
    green = PatternFill("solid", fgColor="11694D")
    overdue_fill = PatternFill("solid", fgColor="BF1F27")
    white = Font(name="Aptos", size=10, color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = green
        cell.font = white
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 31
    for item in rows:
        emission = date.fromisoformat(str(item["dataEmissao"]))
        due = date.fromisoformat(str(item["previsaoChegada"]))
        quantity = Decimal(str(item["quantidade"]))
        sheet.append((emission, _excel_text(item.get("numeroNFe")), due,
                      _excel_text(item["emitente"]), _excel_text(item["ean"]),
                      _excel_text(item["produto"]), float(quantity)))
        line = sheet.max_row
        sheet.cell(line, 1).number_format = "DD/MM/YYYY"
        sheet.cell(line, 2).number_format = "@"
        sheet.cell(line, 3).number_format = "DD/MM/YYYY"
        sheet.cell(line, 5).number_format = "@"
        sheet.cell(line, 7).number_format = "#,##0.####"
        if item.get("atrasado"):
            cell = sheet.cell(line, 3)
            cell.fill = overdue_fill
            cell.font = white
            cell.comment = None
        for col in (1, 3, 7):
            sheet.cell(line, col).alignment = Alignment(vertical="center")
    for column, width in {"A": 20, "B": 16, "C": 24, "D": 42, "E": 20, "F": 62, "G": 17}.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:G{sheet.max_row}"
    data = io.BytesIO()
    workbook.save(data)
    return data.getvalue()


@router.get("/industrias/transito/excel")
async def transit_excel(
    laboratorio: str | None = Query(default=None),
    busca: str = Query(default="", max_length=180),
    session: str | None = Cookie(default=None, alias=settings.cookie_name),
):
    profile = await _industry_profile(session, require_password_changed=True)
    selected = _choose_lab(profile, laboratorio)
    if _is_all_labs_request(selected) and not _buyer_all_labs(profile):
        raise HTTPException(403, "Visão consolidada não autorizada.")
    try:
        payload, _ = await cache_get(modulo=MODULE, settings=settings)
    except CacheReadError as exc:
        raise HTTPException(503, "Base de Trânsito indisponível para exportação.") from exc
    if payload.get("schema") != SCHEMA or not isinstance(payload.get("linhas"), list):
        raise HTTPException(503, "Fotografia de Trânsito inválida.")
    records = _scope(payload["linhas"], selected)
    needle = busca.strip().casefold()
    if needle:
        records = [r for r in records if any(
            needle in str(r.get(column) or "").casefold()
            for column in ("numeroNFe", "emitente", "ean", "produto")
        )]
    excel = _build_transit_excel(records)
    return StreamingResponse(
        io.BytesIO(excel),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="DISMEPE_ONE_TRANSITO.xlsx"',
                 "Cache-Control": "no-store", "X-Content-Type-Options": "nosniff"},
    )


@router.post("/admin/industrias/transito/importar")
async def transit_import(
    request: Request,
    session: str | None = Cookie(default=None, alias=settings.cookie_name),
):
    profile = _strict_admin_profile(session)
    # Exigir mesma origem impede envio autenticado a partir de outra página.
    origin = request.headers.get("origin")
    if origin and origin.rstrip("/") != str(request.base_url).rstrip("/"):
        raise HTTPException(403, "Origem da importação não autorizada.")
    if "application/zip" not in request.headers.get("content-type", "").lower():
        raise HTTPException(415, "Envie ZIP em application/zip.")
    if int(request.headers.get("content-length") or 0) > 10 * 1024 * 1024:
        raise HTTPException(413, "ZIP excede o limite de 10 MB.")
    raw = await request.body()
    lines, stats = _decode_zip(raw)
    from datetime import timezone
    stamp = datetime.now(timezone.utc).isoformat()
    payload = {"schema": SCHEMA, "atualizadoEm": stamp, "linhas": lines}
    try:
        await _edge_admin_write("CACHE_SET", {
            "modulo": MODULE, "payload": payload, "nome": "Trânsito — Portal Indústrias",
            "tamanho": len(lines), "versao": SCHEMA,
            "atualizado_por": str(profile.get("usuario") or ""),
        })
        saved, _ = await cache_get(modulo=MODULE, settings=settings)
        if saved.get("atualizadoEm") != stamp or saved.get("schema") != SCHEMA:
            raise RuntimeError("Gravação não confirmada.")
    except Exception as exc:
        raise HTTPException(503, "Não foi possível confirmar a publicação do Trânsito.") from exc
    return {"sucesso": True, **stats, "atualizadoEm": stamp}
