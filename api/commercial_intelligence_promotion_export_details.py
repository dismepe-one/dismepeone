from __future__ import annotations

import calendar
import io
from datetime import date, datetime
from typing import Any
from zoneinfo import ZoneInfo
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image, LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle


_TZ = ZoneInfo("America/Recife")
_INSTALLED = False


def _money(value: Any) -> str:
    number = float(value or 0)
    return f"R$ {number:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _parse_expiry(value: Any) -> date | None:
    raw = str(value or "").strip()
    if not raw:
        return None
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d/%m/%y"):
        try:
            return datetime.strptime(raw[:10], fmt).date()
        except ValueError:
            pass
    for fmt in ("%m/%Y", "%m-%Y"):
        try:
            dt = datetime.strptime(raw[:7], fmt)
            return date(dt.year, dt.month, calendar.monthrange(dt.year, dt.month)[1])
        except ValueError:
            pass
    return None


def _limit_12_months(today: date) -> date:
    year = today.year + 1
    month = today.month
    day = min(today.day, calendar.monthrange(year, month)[1])
    return date(year, month, day)


def _near_expiry(item: dict[str, Any]) -> bool:
    if bool(item.get("vencimentoProximo")):
        return True
    expiry = _parse_expiry(item.get("vencimento"))
    if expiry is None:
        return False
    today = datetime.now(_TZ).date()
    return today <= expiry <= _limit_12_months(today)


def install_commercial_intelligence_promotion_export_details() -> None:
    global _INSTALLED
    if _INSTALLED:
        return

    from . import commercial_intelligence_promotions as promo

    def excel_bytes(name: str, items: list[dict[str, Any]]) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "Promoção"
        headers = [
            "Código", "Descrição", "Fornecedor", "Custo Médio",
            "Preço Promocional", "Markup Promoção (%)", "Validade",
        ]
        ws["A1"] = name
        ws["A1"].font = Font(size=15, bold=True, color="075B49")
        ws["A2"] = "Gerado em " + datetime.now(_TZ).strftime("%d/%m/%Y %H:%M")
        ws["A2"].font = Font(size=9, color="60746F")
        header_row = 4
        for col, value in enumerate(headers, 1):
            cell = ws.cell(header_row, col, value)
            cell.fill = PatternFill("solid", fgColor="075B49")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")

        for row_index, item in enumerate(items, header_row + 1):
            values = [
                str(item.get("codigo") or ""),
                str(item.get("produto") or ""),
                str(item.get("fornecedor") or ""),
                float(item.get("custo") or 0),
                float(item.get("precoPromocional") or 0),
                float(item.get("markupPromocao") or 0),
                str(item.get("vencimento") or ""),
            ]
            for col, value in enumerate(values, 1):
                ws.cell(row_index, col, value)
            ws.cell(row_index, 4).number_format = 'R$ #,##0.00'
            ws.cell(row_index, 5).number_format = 'R$ #,##0.00'
            ws.cell(row_index, 6).number_format = '0.00"%"'
            if _near_expiry(item) and values[6]:
                ws.cell(row_index, 7).font = Font(color="C00000", bold=True)

        ws.freeze_panes = "A5"
        ws.auto_filter.ref = f"A4:G{max(4, ws.max_row)}"
        widths = (14, 44, 30, 16, 20, 22, 16)
        for index, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(index)].width = width
        for row in ws.iter_rows(min_row=header_row + 1):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    def pdf_bytes(name: str, items: list[dict[str, Any]]) -> bytes:
        output = io.BytesIO()
        doc = SimpleDocTemplate(
            output,
            pagesize=landscape(A4),
            leftMargin=9 * mm,
            rightMargin=9 * mm,
            topMargin=9 * mm,
            bottomMargin=10 * mm,
            title=name,
        )
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "promo-title-details",
            parent=styles["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=14,
            leading=16,
            textColor=colors.HexColor("#075B49"),
            spaceAfter=2 * mm,
        )
        cell_style = ParagraphStyle(
            "promo-cell-details",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=6.2,
            leading=7.4,
            textColor=colors.HexColor("#17332C"),
        )
        note_style = ParagraphStyle(
            "promo-note-details",
            parent=cell_style,
            fontSize=6.5,
            leading=8,
            textColor=colors.HexColor("#60746F"),
        )

        story: list[Any] = []
        logo = promo._logo_bytes()
        if logo:
            story.append(Image(io.BytesIO(logo), width=48 * mm, height=17 * mm))
        story.append(Paragraph(escape(name), title_style))
        story.append(Paragraph("Gerado em " + datetime.now(_TZ).strftime("%d/%m/%Y %H:%M"), note_style))
        story.append(Spacer(1, 3 * mm))

        data: list[list[Any]] = [[
            "Código", "Descrição", "Fornecedor", "Custo Médio", "Preço Promocional", "Markup Promoção"
        ]]
        for item in items:
            data.append([
                str(item.get("codigo") or ""),
                Paragraph(escape(str(item.get("produto") or "")), cell_style),
                Paragraph(escape(str(item.get("fornecedor") or "")), cell_style),
                _money(item.get("custo")),
                _money(item.get("precoPromocional")),
                f"{float(item.get('markupPromocao') or 0):.2f}%".replace(".", ","),
            ])

        table = LongTable(
            data,
            repeatRows=1,
            colWidths=[18 * mm, 62 * mm, 42 * mm, 28 * mm, 32 * mm, 30 * mm],
        )
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#075B49")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
            ("FONTSIZE", (0, 0), (-1, -1), 6.3),
            ("GRID", (0, 0), (-1, -1), .3, colors.HexColor("#D8E5E1")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7FAF9")]),
            ("ALIGN", (3, 1), (-1, -1), "RIGHT"),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(table)
        story.append(Spacer(1, 3 * mm))
        story.append(Paragraph(
            "Este rascunho será excluído automaticamente em até 24 horas após a exportação.",
            note_style,
        ))
        doc.build(story)
        return output.getvalue()

    promo._excel_bytes = excel_bytes
    promo._pdf_bytes = pdf_bytes
    _INSTALLED = True
