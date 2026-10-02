from __future__ import annotations

import io
import os
import re
from datetime import datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

import httpx
from fastapi import Cookie, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image, LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

from .config import get_settings
from .security import normalizar


settings = get_settings()
_TZ = ZoneInfo("America/Recife")
ROOT = Path(__file__).resolve().parents[1]
INDUSTRIES_FILE = ROOT / "frontend" / "industries.html"
_INSTALLED = False


class PromotionItem(BaseModel):
    codigo: str
    produto: str = ""
    fornecedor: str = ""
    custo: float = 0.0
    precoPromocional: float = Field(default=0.0, ge=0)
    vencimento: str = ""
    vencimentoProximo: bool = False


class PromotionSave(BaseModel):
    id: str | None = None
    name: str = "Nova promoção"
    items: list[PromotionItem] = Field(default_factory=list)


def _owner(session: str | None) -> str:
    from . import commercial_intelligence as ci

    profile = ci._profile(session)
    owner = normalizar(profile.get("usuario") or profile.get("sub") or profile.get("nome") or "")
    if not owner:
        raise HTTPException(status_code=401, detail="Usuário da sessão não identificado.")
    return owner


def _secret() -> str:
    value = str(os.getenv("DISMEPE_COMMERCIAL_PROMO_SECRET") or "").strip()
    if not value:
        raise HTTPException(status_code=503, detail="Armazenamento temporário de promoções indisponível.")
    return value


async def _rpc(action: str, owner: str, payload: dict[str, Any] | None = None) -> dict[str, Any]:
    endpoint = settings.supabase_url.rstrip("/") + "/rest/v1/rpc/dismepe_commercial_promo_api"
    key = settings.supabase_publishable_key
    headers = {
        "apikey": key,
        "authorization": f"Bearer {key}",
        "content-type": "application/json",
        "accept": "application/json",
    }
    body = {
        "p_secret": _secret(),
        "p_action": action,
        "p_owner": owner,
        "p_payload": payload or {},
    }
    try:
        async with httpx.AsyncClient(timeout=max(10.0, settings.request_timeout_seconds)) as client:
            response = await client.post(endpoint, headers=headers, json=body)
    except (httpx.TimeoutException, httpx.NetworkError) as exc:
        raise HTTPException(status_code=503, detail="O armazenamento de promoções não respondeu.") from exc

    try:
        data = response.json()
    except ValueError as exc:
        raise HTTPException(status_code=503, detail=f"Resposta inválida do armazenamento (HTTP {response.status_code}).") from exc

    if response.status_code < 200 or response.status_code >= 300:
        message = ""
        if isinstance(data, dict):
            message = str(data.get("message") or data.get("details") or data.get("hint") or "")
        if "not_found" in message:
            raise HTTPException(status_code=404, detail="Promoção não encontrada.")
        raise HTTPException(status_code=503, detail=message or f"Falha ao acessar promoções (HTTP {response.status_code}).")
    if not isinstance(data, dict):
        raise HTTPException(status_code=503, detail="Formato inválido do armazenamento de promoções.")
    return data


def _normalize_items(items: list[PromotionItem]) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    seen: set[str] = set()
    for source in items[:1000]:
        code = str(source.codigo or "").strip()
        if not code or code in seen:
            continue
        seen.add(code)
        cost = max(0.0, float(source.custo or 0))
        price = max(0.0, float(source.precoPromocional or 0))
        markup = ((price / cost) - 1.0) * 100.0 if cost > 0 else 0.0
        result.append({
            "codigo": code,
            "produto": str(source.produto or "").strip(),
            "fornecedor": str(source.fornecedor or "").strip(),
            "custo": round(cost, 4),
            "precoPromocional": round(price, 2),
            "markupPromocao": round(markup, 2),
            "vencimento": str(source.vencimento or "").strip(),
            "vencimentoProximo": bool(source.vencimentoProximo),
        })
    return result


def _logo_bytes() -> bytes | None:
    try:
        text = INDUSTRIES_FILE.read_text(encoding="utf-8")
    except Exception:
        return None
    match = re.search(r'<div class="brand"><img src="data:image/png;base64,([^"]+)"', text)
    if not match:
        return None
    import base64
    try:
        return base64.b64decode(match.group(1))
    except Exception:
        return None


def _excel_bytes(name: str, items: list[dict[str, Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Promoção"
    headers = ["Código", "Preço Promocional", "Markup Promoção (%)"]
    ws["A1"] = name
    ws["A1"].font = Font(size=15, bold=True, color="075B49")
    ws["A2"] = "Gerado em " + datetime.now(_TZ).strftime("%d/%m/%Y %H:%M")
    ws["A2"].font = Font(size=9, color="60746F")
    header_row = 4
    for col, value in enumerate(headers, 1):
        cell = ws.cell(header_row, col, value)
        cell.fill = PatternFill("solid", fgColor="075B49")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    for row_index, item in enumerate(items, header_row + 1):
        ws.cell(row_index, 1, str(item.get("codigo") or ""))
        ws.cell(row_index, 2, float(item.get("precoPromocional") or 0))
        ws.cell(row_index, 3, float(item.get("markupPromocao") or 0))
        ws.cell(row_index, 2).number_format = 'R$ #,##0.00'
        ws.cell(row_index, 3).number_format = '0.00"%"'
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:C{max(4, ws.max_row)}"
    for index, width in enumerate((14, 20, 22), 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def _pdf_bytes(name: str, items: list[dict[str, Any]]) -> bytes:
    output = io.BytesIO()
    doc = SimpleDocTemplate(
        output,
        pagesize=A4,
        leftMargin=13 * mm,
        rightMargin=13 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=name,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "promo-title",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=15,
        leading=18,
        textColor=colors.HexColor("#075B49"),
        spaceAfter=2 * mm,
    )
    note_style = ParagraphStyle(
        "promo-note",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=7.5,
        leading=9,
        textColor=colors.HexColor("#60746F"),
    )
    story: list[Any] = []
    logo = _logo_bytes()
    if logo:
        story.append(Image(io.BytesIO(logo), width=48 * mm, height=17 * mm))
    story.append(Paragraph(name, title_style))
    story.append(Paragraph("Gerado em " + datetime.now(_TZ).strftime("%d/%m/%Y %H:%M"), note_style))
    story.append(Spacer(1, 4 * mm))

    data: list[list[Any]] = [["Código", "Preço Promocional", "Markup Promoção"]]
    for item in items:
        price = float(item.get("precoPromocional") or 0)
        markup = float(item.get("markupPromocao") or 0)
        data.append([
            str(item.get("codigo") or ""),
            f"R$ {price:,.2f}".replace(",", "X").replace(".", ",").replace("X", "."),
            f"{markup:,.2f}%".replace(",", "X").replace(".", ",").replace("X", "."),
        ])
    table = LongTable(data, repeatRows=1, colWidths=[48 * mm, 58 * mm, 58 * mm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#075B49")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), .35, colors.HexColor("#D8E5E1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7FAF9")]),
        ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(table)
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph("Este rascunho será excluído automaticamente em até 24 horas após a exportação.", note_style))
    doc.build(story)
    return output.getvalue()


def install_commercial_intelligence_promotions(app: Any) -> None:
    global _INSTALLED
    if _INSTALLED:
        return
    _INSTALLED = True

    async def list_promotions(session: str | None = Cookie(default=None, alias=settings.cookie_name)):
        owner = _owner(session)
        return await _rpc("LIST", owner)

    async def save_promotion(body: PromotionSave, session: str | None = Cookie(default=None, alias=settings.cookie_name)):
        owner = _owner(session)
        name = str(body.name or "").strip()[:120] or "Nova promoção"
        payload = {"id": body.id, "name": name, "items": _normalize_items(body.items)}
        return await _rpc("UPSERT", owner, payload)

    async def delete_promotion(promotion_id: str, session: str | None = Cookie(default=None, alias=settings.cookie_name)):
        owner = _owner(session)
        return await _rpc("DELETE", owner, {"id": promotion_id})

    async def export_promotion_xlsx(promotion_id: str, session: str | None = Cookie(default=None, alias=settings.cookie_name)):
        owner = _owner(session)
        data = await _rpc("GET", owner, {"id": promotion_id})
        promo = data.get("promocao") if isinstance(data.get("promocao"), dict) else {}
        items = promo.get("items") if isinstance(promo.get("items"), list) else []
        body = _excel_bytes(str(promo.get("name") or "Promoção"), items)
        await _rpc("MARK_EXPORT", owner, {"id": promotion_id})
        filename = "promocao_" + datetime.now(_TZ).strftime("%Y%m%d_%H%M") + ".xlsx"
        return StreamingResponse(io.BytesIO(body), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{filename}"', "Cache-Control": "no-store"})

    async def export_promotion_pdf(promotion_id: str, session: str | None = Cookie(default=None, alias=settings.cookie_name)):
        owner = _owner(session)
        data = await _rpc("GET", owner, {"id": promotion_id})
        promo = data.get("promocao") if isinstance(data.get("promocao"), dict) else {}
        items = promo.get("items") if isinstance(promo.get("items"), list) else []
        body = _pdf_bytes(str(promo.get("name") or "Promoção"), items)
        await _rpc("MARK_EXPORT", owner, {"id": promotion_id})
        filename = "promocao_" + datetime.now(_TZ).strftime("%Y%m%d_%H%M") + ".pdf"
        return StreamingResponse(io.BytesIO(body), media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="{filename}"', "Cache-Control": "no-store"})

    app.add_api_route("/data/inteligencia-comercial/promocoes", list_promotions, methods=["GET"])
    app.add_api_route("/data/inteligencia-comercial/promocoes", save_promotion, methods=["POST"])
    app.add_api_route("/data/inteligencia-comercial/promocoes/{promotion_id}", delete_promotion, methods=["DELETE"])
    app.add_api_route("/data/inteligencia-comercial/promocoes/{promotion_id}/export.xlsx", export_promotion_xlsx, methods=["GET"])
    app.add_api_route("/data/inteligencia-comercial/promocoes/{promotion_id}/export.pdf", export_promotion_pdf, methods=["GET"])
